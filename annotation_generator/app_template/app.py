import os, io, csv, json
from pathlib import Path
from flask import Flask, render_template, request, jsonify, session, redirect, url_for, Response
from models import db, Expert, Assignment, ItemReview, TextAnnotation

BASE = Path(__file__).resolve().parent
CONFIG_PATH = BASE / "config.json"
DATA_DIR = BASE / "data"

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "annotation-tool-dev-key")
app.config["SQLALCHEMY_DATABASE_URI"] = os.environ.get(
    "DATABASE_URL", "sqlite:///annotations.db"
).replace("postgres://", "postgresql://", 1)
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024  # 100MB upload limit

db.init_app(app)

POST_IDS, POSTS, COMMENTS, MODELS = [], {}, {}, []
CONFIG = {}


def load_config():
    global CONFIG
    if CONFIG_PATH.exists():
        CONFIG.update(json.loads(CONFIG_PATH.read_text()))
    return bool(CONFIG)


def init_data():
    global POST_IDS, POSTS, COMMENTS, MODELS
    if CONFIG:
        from load_data import load_all
        POST_IDS, POSTS, COMMENTS, MODELS = load_all()


def is_setup_done():
    return CONFIG_PATH.exists() and (DATA_DIR / "data.csv").exists()


@app.context_processor
def inject_config():
    return {"app_config": CONFIG}


@app.before_request
def check_setup():
    if not is_setup_done() and request.endpoint not in ("setup", "setup_preview", "setup_finish", "static"):
        return redirect(url_for("setup"))


# ── Setup ─────────────────────────────────────────────

@app.route("/setup", methods=["GET"])
def setup():
    if is_setup_done():
        return redirect(url_for("dashboard"))
    return render_template("setup.html")


@app.route("/setup/preview", methods=["POST"])
def setup_preview():
    f = request.files.get("datafile")
    if not f or not f.filename:
        return jsonify({"error": "No file uploaded"}), 400

    ext = f.filename.rsplit(".", 1)[-1].lower()
    if ext == "csv":
        content = f.read().decode("utf-8", errors="replace")
        f.seek(0)
        reader = csv.DictReader(io.StringIO(content))
        headers = reader.fieldnames or []
        sample = []
        for i, row in enumerate(reader):
            if i >= 3:
                break
            sample.append(row)
    elif ext in ("xlsx", "xls"):
        try:
            import openpyxl
        except ImportError:
            return jsonify({"error": "Excel support requires openpyxl (pip install openpyxl)"}), 400
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        f.seek(0)
        if not rows:
            return jsonify({"error": "Empty file"}), 400
        headers = [str(h or "").strip() for h in rows[0]]
        sample = [{h: str(v or "") for h, v in zip(headers, row)} for row in rows[1:4]]
    else:
        return jsonify({"error": f"Unsupported file type: .{ext}"}), 400

    # Auto-detect column names
    auto = {}
    for h in headers:
        hl = h.lower().replace(" ", "_")
        if "post_id" in hl or hl == "id":
            auto.setdefault("post_id", h)
        elif "post_title" in hl or hl == "title":
            auto.setdefault("post_title", h)
        elif "post_body" in hl or hl in ("body", "selftext"):
            auto.setdefault("post_body", h)
        elif "comment_id" in hl:
            auto.setdefault("comment_id", h)
        elif "comment_body" in hl or hl in ("comment", "comment_text"):
            auto.setdefault("comment_body", h)
        elif any(k in hl for k in ("theme", "topic", "label", "tag")):
            auto.setdefault("themes", h)

    return jsonify({"headers": headers, "sample": sample, "auto_map": auto})


@app.route("/setup/finish", methods=["POST"])
def setup_finish():
    f = request.files.get("datafile")
    if not f or not f.filename:
        return jsonify({"error": "No file uploaded"}), 400

    col_post_id = request.form.get("col_post_id", "").strip()
    col_comment_body = request.form.get("col_comment_body", "").strip()
    if not col_post_id or not col_comment_body:
        return jsonify({"error": "Post ID and Comment Body are required"}), 400

    # Convert Excel to CSV if needed
    ext = f.filename.rsplit(".", 1)[-1].lower()
    if ext == "csv":
        csv_content = f.read().decode("utf-8", errors="replace")
    elif ext in ("xlsx", "xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        headers = [str(h or "").strip() for h in rows[0]]
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        for row in rows[1:]:
            writer.writerow({h: str(v or "") for h, v in zip(headers, row)})
        csv_content = output.getvalue()
    else:
        return jsonify({"error": f"Unsupported: .{ext}"}), 400

    config = {
        "app_name": request.form.get("app_name", "Annotation Tool").strip(),
        "admin_password": request.form.get("admin_password", "admin123").strip(),
        "link_template": request.form.get("link_template", "").strip(),
        "instructions": request.form.get("instructions", "").strip(),
        "data_file": "data.csv",
        "columns": {
            "post_id": col_post_id,
            "post_title": request.form.get("col_post_title", "").strip(),
            "post_body": request.form.get("col_post_body", "").strip(),
            "comment_id": request.form.get("col_comment_id", "").strip(),
            "comment_body": col_comment_body,
            "themes": request.form.get("col_themes", "").strip(),
        },
    }

    # Save config and data
    DATA_DIR.mkdir(exist_ok=True)
    CONFIG_PATH.write_text(json.dumps(config, indent=2))
    (DATA_DIR / "data.csv").write_text(csv_content, encoding="utf-8")

    # Initialize
    load_config()
    seed_users()
    init_data()

    return jsonify({"ok": True})


# ── Auth ──────────────────────────────────────────────

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        if not username or not password:
            return render_template("login.html", error="Username and password required")
        expert = Expert.query.filter_by(username=username).first()
        if not expert:
            return render_template("login.html", error="Invalid username or password")
        if not expert.password_hash:
            expert.set_password(password)
            db.session.commit()
        elif not expert.check_password(password):
            return render_template("login.html", error="Invalid username or password")
        session["expert_id"] = expert.id
        session["username"] = expert.username
        return redirect(url_for("dashboard"))
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        confirm = request.form.get("confirm_password", "").strip()
        if not username or not password:
            return render_template("register.html", error="Username and password are required")
        if len(username) < 3:
            return render_template("register.html", error="Username must be at least 3 characters")
        if len(password) < 6:
            return render_template("register.html", error="Password must be at least 6 characters")
        if password != confirm:
            return render_template("register.html", error="Passwords do not match")
        if Expert.query.filter_by(username=username).first():
            return render_template("register.html", error="Username already taken")
        new_expert = Expert(username=username)
        new_expert.set_password(password)
        db.session.add(new_expert)
        db.session.commit()
        session["expert_id"] = new_expert.id
        session["username"] = new_expert.username
        return redirect(url_for("dashboard"))
    return render_template("register.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


def get_expert():
    eid = session.get("expert_id")
    if not eid:
        return None
    return Expert.query.get(eid)


# ── Dashboard ─────────────────────────────────────────

@app.route("/")
def dashboard():
    expert = get_expert()
    if not expert:
        return redirect(url_for("login"))

    search = request.args.get("search", "").strip()
    annotator_filter = request.args.get("annotator", "").strip()
    model_name = "comment_annotations"
    is_admin = expert.username == "admin"

    if not is_admin:
        assigned_ids = {a.post_id for a in Assignment.query.filter_by(expert_id=expert.id).all()}
        if not assigned_ids:
            return render_template("dashboard.html", posts=[], search=search,
                                   username=session.get("username"), is_admin=False,
                                   annotator_filter="", all_annotator_names=[],
                                   total_comments=0, total_annotated=0)

    # Batch DB queries
    if is_admin:
        annot_counts = dict(
            db.session.query(TextAnnotation.post_id, db.func.count(db.func.distinct(TextAnnotation.item_index)))
            .filter_by(model_name=model_name).group_by(TextAnnotation.post_id).all()
        )
        assigned_rows = db.session.query(Assignment.post_id, Expert.username) \
            .join(Expert, Expert.id == Assignment.expert_id).all()
        assigned_map = {}
        for pid_r, uname in assigned_rows:
            assigned_map.setdefault(pid_r, set()).add(uname)
        coded_rows = db.session.query(TextAnnotation.post_id, Expert.username) \
            .join(Expert, Expert.id == TextAnnotation.expert_id) \
            .filter(TextAnnotation.model_name == model_name).distinct().all()
        coded_map = {}
        for pid_r, uname in coded_rows:
            coded_map.setdefault(pid_r, set()).add(uname)
    else:
        annot_counts = dict(
            db.session.query(TextAnnotation.post_id, db.func.count(db.func.distinct(TextAnnotation.item_index)))
            .filter_by(expert_id=expert.id, model_name=model_name).group_by(TextAnnotation.post_id).all()
        )

    posts_list = []
    for pid in POST_IDS:
        if not is_admin and pid not in assigned_ids:
            continue
        key = (pid, model_name)
        post = POSTS.get(key)
        if not post:
            continue
        if search and search.lower() not in (post["title"] or "").lower() and search.lower() not in pid.lower():
            continue
        annotated_comments = annot_counts.get(pid, 0)
        assigned_names = sorted(assigned_map.get(pid, set()) | coded_map.get(pid, set())) if is_admin else []
        posts_list.append({
            "post_id": pid, "title": post["title"], "class_label": post["class_label"],
            "num_comments": len(post["advice"]), "annotated_comments": annotated_comments,
            "assigned_names": assigned_names, "link": post["link"],
        })

    all_annotator_names = sorted({name for p in posts_list for name in p.get("assigned_names", [])}) if is_admin else []
    if is_admin and annotator_filter:
        posts_list = [p for p in posts_list if annotator_filter in p.get("assigned_names", [])]
    posts_list.sort(key=lambda p: -p["num_comments"])

    total_comments = sum(p["num_comments"] for p in posts_list)
    total_annotated = sum(p["annotated_comments"] for p in posts_list)

    return render_template("dashboard.html", posts=posts_list, search=search,
                           annotator_filter=annotator_filter, all_annotator_names=all_annotator_names,
                           username=session.get("username"), is_admin=is_admin,
                           total_comments=total_comments, total_annotated=total_annotated)


# ── Review Page ───────────────────────────────────────

@app.route("/review/<post_id>")
def review(post_id):
    expert = get_expert()
    if not expert:
        return redirect(url_for("login"))
    if expert.username != "admin":
        assigned_ids = {a.post_id for a in Assignment.query.filter_by(expert_id=expert.id).all()}
        if post_id not in assigned_ids:
            return "Not assigned to this post", 403

    model_name = "comment_annotations"
    post = POSTS.get((post_id, model_name))
    if not post:
        return "Post not found", 404
    comment_data = COMMENTS.get(post_id, {})

    existing_annotations = []
    for a in TextAnnotation.query.filter_by(expert_id=expert.id, post_id=post_id, model_name=model_name).all():
        existing_annotations.append({
            "id": a.id, "section": a.section, "item_index": a.item_index,
            "start": a.start_offset, "end": a.end_offset, "text": a.highlighted_text,
            "annotation": a.annotation_text, "verdict": a.verdict,
            "harm_verdict": a.harm_verdict or "", "factual_reasoning": a.factual_reasoning or "",
            "harm_reasoning": a.harm_reasoning or "", "is_gt_span": a.is_gt_span or False,
        })

    if expert.username != "admin":
        nav_ids = [pid for pid in POST_IDS if pid in assigned_ids]
    else:
        nav_ids = POST_IDS
    try:
        idx = nav_ids.index(post_id)
    except ValueError:
        idx = 0
    prev_id = nav_ids[idx - 1] if idx > 0 else None
    next_id = nav_ids[idx + 1] if idx < len(nav_ids) - 1 else None

    return render_template("review.html", post=post, comment_data=comment_data,
                           existing_annotations=existing_annotations, current_model=model_name,
                           prev_id=prev_id, next_id=next_id, username=session.get("username"))


# ── API ───────────────────────────────────────────────

@app.route("/api/annotation/<post_id>/save", methods=["POST"])
def save_annotation(post_id):
    expert = get_expert()
    if not expert:
        return jsonify({"error": "Not logged in"}), 401
    data = request.json
    section = data["section"]
    item_index = data.get("item_index", 0)
    start, end = data["start"], data["end"]

    overlapping = TextAnnotation.query.filter_by(
        expert_id=expert.id, post_id=post_id, model_name=data["model_name"],
        section=section, item_index=item_index,
    ).filter(TextAnnotation.start_offset < end, TextAnnotation.end_offset > start).all()
    removed_ids = [a.id for a in overlapping]
    for a in overlapping:
        db.session.delete(a)

    annot = TextAnnotation(
        expert_id=expert.id, post_id=post_id, model_name=data["model_name"],
        section=section, item_index=item_index, start_offset=start, end_offset=end,
        highlighted_text=data["text"], annotation_text=data.get("annotation", ""),
        verdict=data.get("verdict"), harm_verdict=data.get("harm_verdict"),
        factual_reasoning=data.get("factual_reasoning", ""),
        harm_reasoning=data.get("harm_reasoning", ""),
        is_gt_span=data.get("is_gt_span", False),
    )
    db.session.add(annot)
    db.session.commit()
    return jsonify({"ok": True, "id": annot.id, "removed_ids": removed_ids})


@app.route("/api/annotation/<int:annot_id>/delete", methods=["POST"])
def delete_annotation(annot_id):
    expert = get_expert()
    if not expert:
        return jsonify({"error": "Not logged in"}), 401
    annot = TextAnnotation.query.get(annot_id)
    if annot and annot.expert_id == expert.id:
        db.session.delete(annot)
        db.session.commit()
    return jsonify({"ok": True})


# ── Admin ─────────────────────────────────────────────

@app.route("/admin/review/<post_id>")
def admin_review(post_id):
    expert = get_expert()
    if not expert or expert.username != "admin":
        return redirect(url_for("login"))
    model_name = "comment_annotations"
    post = POSTS.get((post_id, model_name))
    if not post:
        return "Post not found", 404
    comment_data = COMMENTS.get(post_id, {})

    all_annotations = []
    annotations_by_expert = {}
    for a in TextAnnotation.query.filter_by(post_id=post_id, model_name=model_name).all():
        expert_obj = Expert.query.get(a.expert_id)
        annot = {
            "id": a.id, "expert_id": a.expert_id,
            "expert_name": expert_obj.username if expert_obj else "unknown",
            "section": a.section, "item_index": a.item_index,
            "start": a.start_offset, "end": a.end_offset, "text": a.highlighted_text,
            "annotation": a.annotation_text, "verdict": a.verdict,
            "harm_verdict": a.harm_verdict or "", "factual_reasoning": a.factual_reasoning or "",
            "harm_reasoning": a.harm_reasoning or "",
        }
        all_annotations.append(annot)
        annotations_by_expert.setdefault(annot["expert_name"], []).append(annot)

    try:
        idx = POST_IDS.index(post_id)
    except ValueError:
        idx = 0
    prev_id = POST_IDS[idx - 1] if idx > 0 else None
    next_id = POST_IDS[idx + 1] if idx < len(POST_IDS) - 1 else None

    return render_template("admin_review.html", post=post, comment_data=comment_data,
                           all_annotations=all_annotations, annotations_by_expert=annotations_by_expert,
                           expert_names=sorted(annotations_by_expert.keys()),
                           current_model=model_name, prev_id=prev_id, next_id=next_id,
                           username=session.get("username"))


@app.route("/admin/export_csv")
def admin_export_csv():
    expert = get_expert()
    if not expert or expert.username != "admin":
        return redirect(url_for("login"))
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["annotation_id", "expert", "post_id", "comment_index",
                     "highlighted_span", "start_offset", "end_offset",
                     "factual_accuracy", "harm_potential",
                     "factual_reasoning", "harm_reasoning", "optional_comment", "created_at"])
    for a in TextAnnotation.query.order_by(TextAnnotation.post_id, TextAnnotation.item_index, TextAnnotation.start_offset).all():
        expert_obj = Expert.query.get(a.expert_id)
        writer.writerow([a.id, expert_obj.username if expert_obj else "unknown", a.post_id,
                         a.item_index, a.highlighted_text, a.start_offset, a.end_offset,
                         a.verdict or "", a.harm_verdict or "", a.factual_reasoning or "",
                         a.harm_reasoning or "", a.annotation_text or "",
                         a.created_at.isoformat() if a.created_at else ""])
    output.seek(0)
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=annotations_export.csv"})


@app.route("/admin", methods=["GET"])
def admin():
    expert = get_expert()
    if not expert or expert.username != "admin":
        return redirect(url_for("login"))
    experts = Expert.query.all()
    assignments = {e.id: [a.post_id for a in Assignment.query.filter_by(expert_id=e.id).all()] for e in experts}

    taken_by = {}
    for e in experts:
        if e.username == "admin":
            continue
        for pid in assignments.get(e.id, []):
            taken_by.setdefault(pid, []).append(e.username)

    post_meta = {}
    for pid in POST_IDS:
        post = POSTS.get((pid, "comment_annotations"))
        if post:
            post_meta[pid] = {"themes": post["class_label"], "num_comments": len(post["advice"])}

    expert_stats = {}
    for e in experts:
        if e.username == "admin":
            continue
        total_comments = sum(len(POSTS.get((pid, "comment_annotations"), {}).get("advice", []))
                            for pid in assignments.get(e.id, []))
        annotated_comments = db.session.query(
            TextAnnotation.post_id, TextAnnotation.item_index
        ).filter_by(expert_id=e.id).distinct().count()
        expert_stats[e.id] = {"total_comments": total_comments,
                              "annotations": TextAnnotation.query.filter_by(expert_id=e.id).count(),
                              "annotated_comments": annotated_comments}

    return render_template("admin.html", experts=experts, assignments=assignments,
                           taken_by=taken_by, post_meta=post_meta, expert_stats=expert_stats,
                           all_post_ids=POST_IDS, username=session.get("username"))


@app.route("/admin/add_expert", methods=["POST"])
def admin_add_expert():
    expert = get_expert()
    if not expert or expert.username != "admin":
        return jsonify({"error": "Unauthorized"}), 403
    data = request.json
    username = data.get("username", "").strip()
    password = data.get("password", "").strip()
    if not username or not password:
        return jsonify({"error": "Username and password required"}), 400
    if Expert.query.filter_by(username=username).first():
        return jsonify({"error": "Username already exists"}), 400
    new_expert = Expert(username=username, password_plain=password)
    new_expert.set_password(password)
    db.session.add(new_expert)
    db.session.commit()
    return jsonify({"ok": True, "id": new_expert.id})


@app.route("/admin/delete_expert", methods=["POST"])
def admin_delete_expert():
    expert = get_expert()
    if not expert or expert.username != "admin":
        return jsonify({"error": "Unauthorized"}), 403
    target = Expert.query.get(request.json.get("expert_id"))
    if not target:
        return jsonify({"error": "Expert not found"}), 404
    if target.username == "admin":
        return jsonify({"error": "Cannot delete admin"}), 400
    TextAnnotation.query.filter_by(expert_id=target.id).delete()
    ItemReview.query.filter_by(expert_id=target.id).delete()
    Assignment.query.filter_by(expert_id=target.id).delete()
    db.session.delete(target)
    db.session.commit()
    return jsonify({"ok": True})


@app.route("/admin/assign", methods=["POST"])
def admin_assign():
    expert = get_expert()
    if not expert or expert.username != "admin":
        return jsonify({"error": "Unauthorized"}), 403
    data = request.json
    expert_id = data.get("expert_id")
    Assignment.query.filter_by(expert_id=expert_id).delete()
    for pid in data.get("post_ids", []):
        if not Assignment.query.filter_by(expert_id=expert_id, post_id=pid).first():
            db.session.add(Assignment(expert_id=expert_id, post_id=pid))
    db.session.commit()
    return jsonify({"ok": True})


@app.route("/admin/db")
def admin_db():
    expert = get_expert()
    if not expert or expert.username != "admin":
        return redirect(url_for("login"))
    from sqlalchemy import text, inspect
    inspector = inspect(db.engine)
    table_names = inspector.get_table_names()
    selected = request.args.get("table", "")
    rows, columns, row_count = [], [], 0
    if selected and selected in table_names:
        row_count = db.session.execute(text(f'SELECT count(*) FROM "{selected}"')).scalar()
        limit = int(request.args.get("limit", 100))
        offset = int(request.args.get("offset", 0))
        result = db.session.execute(text(f'SELECT * FROM "{selected}" ORDER BY 1 DESC LIMIT :lim OFFSET :off'),
                                    {"lim": limit, "off": offset})
        columns = list(result.keys())
        rows = [list(r) for r in result.fetchall()]
    table_counts = {t: db.session.execute(text(f'SELECT count(*) FROM "{t}"')).scalar() for t in table_names}
    return render_template("admin_db.html", tables=table_names, table_counts=table_counts,
                           selected=selected, columns=columns, rows=rows, row_count=row_count,
                           limit=int(request.args.get("limit", 100)), offset=int(request.args.get("offset", 0)))


# ── Startup ───────────────────────────────────────────

def seed_users():
    admin_pw = CONFIG.get("admin_password", "admin123")
    expert = Expert.query.filter_by(username="admin").first()
    if not expert:
        expert = Expert(username="admin")
        db.session.add(expert)
    expert.set_password(admin_pw)
    expert.password_plain = admin_pw
    db.session.commit()


with app.app_context():
    db.create_all()
    if load_config():
        seed_users()
        init_data()

if __name__ == "__main__":
    app.run(debug=True, port=5001, use_reloader=False)
