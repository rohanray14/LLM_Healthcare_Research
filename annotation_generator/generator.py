"""Annotation Tool Generator — upload a dataset, get a deployable annotation app."""

import os, csv, json, shutil, tempfile, io
from pathlib import Path
from flask import Flask, render_template, request, send_file, jsonify

app = Flask(__name__)
app.secret_key = "generator-dev-key"
TEMPLATE_DIR = Path(__file__).resolve().parent / "app_template"


def read_csv_headers(file_storage):
    """Read CSV headers and first few rows from an uploaded file."""
    content = file_storage.read().decode("utf-8", errors="replace")
    file_storage.seek(0)
    reader = csv.DictReader(io.StringIO(content))
    headers = reader.fieldnames or []
    sample_rows = []
    for i, row in enumerate(reader):
        if i >= 3:
            break
        sample_rows.append(row)
    return headers, sample_rows, content


def read_excel_headers(file_storage):
    """Read headers from an Excel file, converting to CSV in memory."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_storage.read()), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    file_storage.seek(0)
    if not rows:
        return [], [], ""
    headers = [str(h or "").strip() for h in rows[0]]
    sample_rows = []
    for row in rows[1:4]:
        sample_rows.append({h: str(v or "") for h, v in zip(headers, row)})
    # Convert to CSV string
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in rows[1:]:
        writer.writerow({h: str(v or "") for h, v in zip(headers, row)})
    return headers, sample_rows, output.getvalue()


@app.route("/", methods=["GET"])
def index():
    return render_template("index.html")


@app.route("/preview", methods=["POST"])
def preview():
    """Upload file and return detected columns for mapping."""
    f = request.files.get("datafile")
    if not f or not f.filename:
        return jsonify({"error": "No file uploaded"}), 400

    ext = f.filename.rsplit(".", 1)[-1].lower()
    if ext == "csv":
        headers, sample_rows, _ = read_csv_headers(f)
    elif ext in ("xlsx", "xls"):
        try:
            headers, sample_rows, _ = read_excel_headers(f)
        except ImportError:
            return jsonify({"error": "Install openpyxl to read Excel files: pip install openpyxl"}), 400
    else:
        return jsonify({"error": f"Unsupported file type: .{ext}"}), 400

    # Auto-detect common column names
    auto_map = {}
    for h in headers:
        hl = h.lower().replace(" ", "_")
        if "post_id" in hl or hl == "id":
            auto_map.setdefault("post_id", h)
        elif "post_title" in hl or hl == "title":
            auto_map.setdefault("post_title", h)
        elif "post_body" in hl or hl == "body" or hl == "selftext":
            auto_map.setdefault("post_body", h)
        elif "comment_id" in hl:
            auto_map.setdefault("comment_id", h)
        elif "comment_body" in hl or hl == "comment" or hl == "comment_text":
            auto_map.setdefault("comment_body", h)
        elif "theme" in hl or "topic" in hl or "label" in hl or "tag" in hl:
            auto_map.setdefault("themes", h)

    return jsonify({"headers": headers, "sample": sample_rows, "auto_map": auto_map})


@app.route("/generate", methods=["POST"])
def generate():
    """Generate the annotation app and return as a zip file."""
    f = request.files.get("datafile")
    if not f or not f.filename:
        return jsonify({"error": "No file uploaded"}), 400

    # Read form config
    app_name = request.form.get("app_name", "Annotation Tool").strip()
    admin_password = request.form.get("admin_password", "admin123").strip()
    link_template = request.form.get("link_template", "").strip()
    instructions = request.form.get("instructions", "").strip()

    col_post_id = request.form.get("col_post_id", "").strip()
    col_post_title = request.form.get("col_post_title", "").strip()
    col_post_body = request.form.get("col_post_body", "").strip()
    col_comment_id = request.form.get("col_comment_id", "").strip()
    col_comment_body = request.form.get("col_comment_body", "").strip()
    col_themes = request.form.get("col_themes", "").strip()

    if not col_post_id or not col_comment_body:
        return jsonify({"error": "Post ID and Comment Body columns are required"}), 400

    # Convert Excel to CSV if needed
    ext = f.filename.rsplit(".", 1)[-1].lower()
    if ext == "csv":
        csv_content = f.read().decode("utf-8", errors="replace")
    elif ext in ("xlsx", "xls"):
        _, _, csv_content = read_excel_headers(f)
        # Re-read for full content
        f.seek(0)
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        headers = [str(h or "").strip() for h in rows[0]]
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        for row in rows[1:]:
            writer.writerow({h: str(v or "") for h, v in zip(headers, row)})
        csv_content = output.getvalue()
    else:
        return jsonify({"error": f"Unsupported: .{ext}"}), 400

    config = {
        "app_name": app_name,
        "admin_password": admin_password,
        "link_template": link_template,
        "instructions": instructions,
        "data_file": "data.csv",
        "columns": {
            "post_id": col_post_id,
            "post_title": col_post_title,
            "post_body": col_post_body,
            "comment_id": col_comment_id,
            "comment_body": col_comment_body,
            "themes": col_themes,
        },
    }

    # Create the app in a temp directory
    tmp = tempfile.mkdtemp()
    app_dir = Path(tmp) / app_name.lower().replace(" ", "_")

    # Copy template
    shutil.copytree(TEMPLATE_DIR, app_dir)

    # Write config
    (app_dir / "config.json").write_text(json.dumps(config, indent=2))

    # Write data
    (app_dir / "data" / "data.csv").write_text(csv_content, encoding="utf-8")

    # Update render.yaml with app name
    render_yaml = app_dir / "render.yaml"
    render_content = render_yaml.read_text()
    render_content = render_content.replace("annotation-tool", app_name.lower().replace(" ", "-"))
    render_yaml.write_text(render_content)

    # Create zip
    zip_path = Path(tmp) / f"{app_dir.name}.zip"
    shutil.make_archive(str(zip_path.with_suffix("")), "zip", tmp, app_dir.name)

    return send_file(str(zip_path), as_attachment=True,
                     download_name=f"{app_dir.name}.zip", mimetype="application/zip")


if __name__ == "__main__":
    app.run(debug=True, port=5050)
